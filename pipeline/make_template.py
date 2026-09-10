#!/usr/bin/env python3
"""Create the risk-register TEMPLATE (generic, no real product).
Two sheets: 'Document Details' (header block) and 'Risk Register' (column headers only).
The pipeline loads this template and appends finding rows, preserving the format.
Run once to (re)generate templates/risk_register_template.xlsx.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "templates", "risk_register_template.xlsx")
ACCENT = "0C8F9B"
thin = Side(style="thin", color="B9C5D2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

COLUMNS = [
    ("Risk ID", 16), ("Category", 12), ("Threat / Finding", 40), ("Asset / Location", 26),
    ("CWE / CVE", 14), ("Severity", 10), ("Likelihood (Pre-RCM)", 16), ("Risk Score", 11),
    ("Risk Level", 14), ("Risk Control Measure", 40), ("Likelihood (Post-RCM)", 16),
    ("Residual Level", 14), ("Acceptability", 16), ("BRA Status", 12), ("Evidence", 24),
]

DOC_DETAILS = [
    ("Document ID", "# DEMO-RR-00x"),
    ("Template ID / Revision", "RR-TEMPLATE and 01"),
    ("Revision No.", "# xx"),
    ("Effective Date", "YYYY-MM-DD"),
    ("Product / System", "[# Sample Application - Version]"),
    ("Prepared by", "[name]"),
    ("Reviewed by", "[name]"),
    ("Purpose", "Security risk register (DEMONSTRATION). Records identified findings, their risk "
                "scoring (likelihood x severity), risk-control measures, residual risk and BRA status. "
                "Sample data only - not a regulatory document."),
]


def main():
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb = openpyxl.Workbook()

    dd = wb.active; dd.title = "Document Details"
    dd["A1"] = "SECURITY RISK REGISTER - TEMPLATE (DEMONSTRATION)"
    dd["A1"].font = Font(bold=True, size=13, color="B23B33"); dd.merge_cells("A1:C1")
    r = 3
    for k, v in DOC_DETAILS:
        dd.cell(r, 1, k).font = Font(bold=True)
        dd.cell(r, 1).fill = PatternFill("solid", fgColor="EEF2F6")
        dd.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
        dd.cell(r, 1).border = BORDER; dd.cell(r, 2).border = BORDER
        r += 1
    dd.column_dimensions["A"].width = 24; dd.column_dimensions["B"].width = 90

    rr = wb.create_sheet("Risk Register")
    rr["A1"] = "DEMONSTRATION - populated automatically by the security pipeline - NOT a regulatory document"
    rr["A1"].font = Font(bold=True, color="B23B33", size=11)
    rr.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    for c, (name, w) in enumerate(COLUMNS, start=1):
        cell = rr.cell(2, c, name)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); cell.border = BORDER
        rr.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    rr.freeze_panes = "A3"
    rr.row_dimensions[2].height = 30

    wb.save(OUT)
    print("wrote template:", OUT, "| columns:", len(COLUMNS))


if __name__ == "__main__":
    main()
