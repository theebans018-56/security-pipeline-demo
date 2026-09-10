#!/usr/bin/env python3
"""
Security compliance pipeline — DEMO with SYNTHETIC DATA ONLY.
Demonstrates: scan -> ledger -> triage/score -> generate register -> verify.
Generic sample application. No real product, no real findings.
"""
import json, os, argparse, datetime, re
HERE = os.path.dirname(os.path.abspath(__file__))
SEV = {"Critical": 5, "High": 4, "Medium": 3, "Low": 2, "Very Low": 1}


def load_findings(path):
    """Synthetic scanner output (dependency + code findings)."""
    return json.load(open(path, encoding="utf-8"))


def load_host_results(path):
    """Results from the SSH compliance test step (PASS/FAIL + observed value)."""
    if path and os.path.exists(path):
        return json.load(open(path, encoding="utf-8"))
    return []


def build_ledger(findings, host_results):
    rows = []
    for i, f in enumerate(findings, start=1):
        rows.append({
            "id": f"FND-{i:04d}", "category": f.get("type", "SCA"),
            "scanner": {"cwe": f.get("cwe", ""), "cve": f.get("cve", ""), "severity": f.get("severity", "Medium"),
                        "location": f.get("location", ""), "observed": f.get("summary", ""), "verdict": "OPEN"},
            "human": {}, "evidence": {"tool": f.get("tool", "scanner"), "ts": now()}}
        )
    for j, h in enumerate(host_results, start=1):
        cat = h.get("category", "COMPLIANCE")           # SRT / TMT / COMPLIANCE
        tid = h.get("test_id", f"CTRL-{j:04d}")
        rows.append({
            "id": tid, "category": cat,
            "scanner": {"cwe": h.get("cwe", ""), "cve": "", "severity": h.get("severity", "Medium"),
                        "location": h.get("command", ""), "observed": h.get("observed", ""), "verdict": h.get("verdict", "PASS")},
            "human": {}, "evidence": {"test_id": tid, "title": h.get("title", ""), "host": h.get("host", ""), "ts": now()}}
        )
    return rows


def triage(rows):
    for r in rows:
        s = r["scanner"]; sev = SEV.get(s["severity"], 3)
        if s["verdict"] == "PASS":                       # a control that PASSED is verified, not a risk
            r["human"] = {"likelihood": "-", "score": 0, "band": "Verified", "bra": "N/A",
                          "disposition": "Verified - control effective", "note": "SRT/TMT PASS"}
            continue
        lik = 3 if s["verdict"] in ("OPEN", "FAIL") else 2
        score = sev * lik
        band = "Low" if score <= 6 else "Moderate" if score <= 12 else "High"
        r["human"] = {"likelihood": {4: "High", 3: "Medium", 2: "Low"}[lik], "score": score,
                      "band": band, "bra": "Applicable" if 7 <= score <= 12 else "N/A",
                      "disposition": "Acceptable (BRA)" if 7 <= score <= 12 else "Acceptable",
                      "note": "failed control re-opened as risk" if s["verdict"] == "FAIL" else "open finding"}
    return rows


TEMPLATE = os.path.join(HERE, "..", "templates", "risk_register_template.xlsx")
FILL = {"Moderate": "F6EBD6", "High": "F5E0DD"}


def write_register(rows, out):
    """Load the register TEMPLATE and populate the 'Risk Register' sheet, preserving its
    Document Details header and column formatting. Also stamps the revision/date."""
    import openpyxl
    from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
    from copy import copy
    thin = Side(style="thin", color="B9C5D2"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if os.path.exists(TEMPLATE):
        wb = openpyxl.load_workbook(TEMPLATE)
    else:  # fallback: regenerate the template on the fly
        import subprocess, sys
        subprocess.run([sys.executable, os.path.join(HERE, "make_template.py")], check=True)
        wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Risk Register"]
    # stamp revision + date in Document Details
    if "Document Details" in wb.sheetnames:
        dd = wb["Document Details"]
        for r in range(1, dd.max_row + 1):
            key = str(dd.cell(r, 1).value or "")
            if key == "Revision No.": dd.cell(r, 2, "01")
            if key == "Effective Date": dd.cell(r, 2, now()[:10])
    # append RISK rows starting after the header (row 2). Passed controls are verified, not risks.
    start = 3
    risks = [r for r in rows if r["human"]["band"] != "Verified"]
    for i, r in enumerate(risks):
        s, h = r["scanner"], r["human"]
        ev = r.get("evidence", {})
        vals = [r["id"], r["category"], str(s["observed"])[:80], s.get("location", ""), s["cve"] or s["cwe"],
                s["severity"], h["likelihood"], h["score"], f"{h['band']} ({h['score']})",
                "Planned on-cycle fix; monitored under vulnerability management.", h["likelihood"],
                h["band"], h["disposition"], h["bra"], ev.get("tool") or ev.get("test_id") or ""]
        row = start + i
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row, c, v); cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
        if h["band"] in FILL:
            ws.cell(row, 12).fill = PatternFill("solid", fgColor=FILL[h["band"]])
            ws.cell(row, 14).fill = PatternFill("solid", fgColor=FILL[h["band"]])
    wb.save(out)


def verify(rows):
    problems = []
    ids = [r["id"] for r in rows]
    if len(ids) != len(set(ids)): problems.append("duplicate IDs")
    for r in rows:
        h = r["human"]
        if h["band"] == "Moderate" and h["bra"] != "Applicable": problems.append(f"{r['id']}: Moderate without BRA")
        if h["band"] == "Low" and h["bra"] == "Applicable": problems.append(f"{r['id']}: Low marked BRA")
        if h["band"] == "Verified" and r["scanner"]["verdict"] != "PASS": problems.append(f"{r['id']}: Verified but not PASS")
    return problems


def now(): return datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%MZ")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--findings", default=os.path.join(HERE, "sample_findings.json"))
    ap.add_argument("--host-results", default=os.path.join(HERE, "..", "host_results.json"))
    ap.add_argument("--compliance", default=os.path.join(HERE, "..", "compliance_results.json"),
                    help="SRT/TMT results JSON from run_compliance.py")
    ap.add_argument("--out-dir", default=HERE)
    a = ap.parse_args()
    print("=== SECURITY COMPLIANCE PIPELINE (demo, synthetic data) ===\n")
    out = os.path.join(a.out_dir, "output"); os.makedirs(out, exist_ok=True)
    controls = load_host_results(a.host_results) + load_host_results(a.compliance)
    rows = triage(build_ledger(load_findings(a.findings), controls))
    led = os.path.join(out, "ledger.jsonl")
    with open(led, "w", encoding="utf-8") as f:
        for r in rows: f.write(json.dumps(r) + "\n")
    print(f"[1-2] scan + ledger  -> {len(rows)} findings -> output/ledger.jsonl")
    for r in rows:
        s, h = r["scanner"], r["human"]
        print(f"      {r['id']:<10} {r['category']:<11} {s['verdict']:<5} sev={s['severity']:<8} score={h['score']:<3} {h['band']:<9} BRA={h['bra']}")
    reg = os.path.join(out, "register.xlsx"); write_register(rows, reg)
    print(f"\n[3-4] triage + document -> output/register.xlsx (filled from templates/risk_register_template.xlsx)")
    # SRT/TMT verification summary
    for cat in ("SRT", "TMT"):
        cr = [r for r in rows if r["category"] == cat]
        if cr:
            p = sum(1 for r in cr if r["scanner"]["verdict"] == "PASS")
            print(f"      {cat}: {p} PASS / {len(cr)-p} FAIL  (failed controls become register risks)")
    problems = verify(rows)
    print("\n[5]   verify -> " + ("PASS" if not problems else "ISSUES: " + "; ".join(problems)))
    if problems: raise SystemExit(1)
    risks = [r for r in rows if r["human"]["band"] != "Verified"]
    bra = [r["id"] for r in risks if r["human"]["bra"] == "Applicable"]
    print(f"\nRESULT: {len(rows)} results ({len(risks)} risks in register, {len(rows)-len(risks)} verified), "
          f"{len(bra)} to BRA ({', '.join(bra) or 'none'}). Synthetic only.")
